import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def generate_presentation_deck(output_path):
    prs = Presentation()
    
    # Slide Dimensions: 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Premium Colors
    bg_color = RGBColor(15, 23, 42)        # Slate-900 (#0F172A)
    card_color = RGBColor(30, 41, 59)      # Slate-800 (#1E293B)
    primary_color = RGBColor(99, 102, 241) # Indigo-500 (#6366F1)
    accent_color = RGBColor(20, 184, 166)  # Teal-500 (#14B8A6)
    white_text = RGBColor(248, 250, 252)   # Slate-50 (#F8FAFC)
    muted_text = RGBColor(148, 163, 184)   # Slate-400 (#94A3B8)
    
    # Helper: Set dark background
    def set_slide_bg(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = bg_color
        
    # Helper: Add Slide Header
    def add_slide_header(slide, title_text):
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.33), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = title_text.upper()
        p.font.name = "Arial"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = primary_color
        
        # Underline accent line
        line = slide.shapes.add_shape(
            1, # Rectangle
            Inches(0.5), Inches(1.1), Inches(2.0), Inches(0.04)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = accent_color
        line.line.color.rgb = accent_color
        line.line.width = Pt(0)
        
    slide_layout = prs.slide_layouts[6] # Blank Layout
    
    # -------------------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # -------------------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide)
    
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.33), Inches(3.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "AURA BI: ENTERPRISE MULTI-AGENT BI PLATFORM"
    p.font.name = "Arial"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = white_text
    
    p2 = tf.add_paragraph()
    p2.text = "Technical Assessment & Portfolio Presentation"
    p2.font.name = "Arial"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = accent_color
    p2.space_before = Pt(8)
    
    p3 = tf.add_paragraph()
    p3.text = "Candidate Implementation | LangGraph Swarms, FastAPI, and React Client"
    p3.font.name = "Arial"
    p3.font.size = Pt(14)
    p3.font.color.rgb = muted_text
    p3.space_before = Pt(45)
    
    # -------------------------------------------------------------------------
    # SLIDE 2: PROJECT OVERVIEW & SCOPE
    # -------------------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide)
    add_slide_header(slide, "Project Overview & Engineering Scope")
    
    # Left Column: The Problem & Value
    left_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(5.8), Inches(4.8))
    ltf = left_box.text_frame
    ltf.word_wrap = True
    
    lp1 = ltf.paragraphs[0]
    lp1.text = "The Business Problem"
    lp1.font.name = "Arial"
    lp1.font.size = Pt(18)
    lp1.font.bold = True
    lp1.font.color.rgb = accent_color
    
    lp2 = ltf.add_paragraph()
    lp2.text = "Traditional corporate profiling and competitor research are highly manual, unstructured, and time-intensive.\n\n" \
               "AURA BI acts as an automated consulting analyst that transforms a single company query into a high-fidelity business profile, SWOT grid, peer benchmarking card, and ROI-driven AI roadmap in seconds."
    lp2.font.name = "Arial"
    lp2.font.size = Pt(13)
    lp2.font.color.rgb = white_text
    lp2.space_before = Pt(10)
    
    # Right Column: Key Implemented Deliverables
    card = slide.shapes.add_shape(1, Inches(6.9), Inches(1.8), Inches(5.9), Inches(4.8))
    card.fill.solid()
    card.fill.fore_color.rgb = card_color
    card.line.color.rgb = primary_color
    card.line.width = Pt(1.5)
    
    rtf = card.text_frame
    rtf.word_wrap = True
    rtf.margin_left = rtf.margin_top = rtf.margin_right = rtf.margin_bottom = Inches(0.3)
    
    rp1 = rtf.paragraphs[0]
    rp1.text = "Key Deliverables"
    rp1.font.name = "Arial"
    rp1.font.size = Pt(18)
    rp1.font.bold = True
    rp1.font.color.rgb = primary_color
    
    bullets = [
        "Dynamic Multi-Agent Swarm (LangGraph nodes)",
        "REST API endpoints with Pydantic validation",
        "Responsive, animated dashboard UI (React 19 & Tailwind)",
        "Widescreen charts (Recharts) & stepper load bars",
        "In-Memory binary file compilers (PDF & PowerPoint)"
    ]
    for b in bullets:
        bp = rtf.add_paragraph()
        bp.text = f"• {b}"
        bp.font.name = "Arial"
        bp.font.size = Pt(13.5)
        bp.font.color.rgb = white_text
        bp.space_before = Pt(10)
        
    # -------------------------------------------------------------------------
    # SLIDE 3: SYSTEM ARCHITECTURE
    # -------------------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide)
    add_slide_header(slide, "System Architecture Design")
    
    # 3 horizontal boxes for stack parts
    widths = Inches(3.8)
    height = Inches(4.8)
    gap = Inches(0.4)
    start_left = Inches(0.5)
    
    stack = [
        ("FRONTEND CLIENT", ["Vite / React 19 / TypeScript", "Tailwind CSS glassmorphism UI", "Framer Motion state steppers", "Recharts dynamic charts", "REST communication via Axios/Fetch"], RGBColor(20, 184, 166)),
        ("BACKEND API SERVER", ["FastAPI & Uvicorn routing", "Asynchronous endpoints", "Pydantic JSON schema parser", "In-memory BytesIO generation", "ReportLab & python-pptx"], RGBColor(99, 102, 241)),
        ("AI DATA ENGINE", ["LangGraph sequential agent swarm", "Tavily LLM-optimized web search", "Google Gemini 1.5 Flash", "Unified LangChain frameworks", "Deterministic fallback models"], RGBColor(168, 85, 247))
    ]
    
    for idx, (title, items, color) in enumerate(stack):
        left = start_left + idx * (widths + gap)
        top = Inches(1.8)
        
        card = slide.shapes.add_shape(1, left, top, widths, height)
        card.fill.solid()
        card.fill.fore_color.rgb = card_color
        card.line.color.rgb = color
        card.line.width = Pt(1.5)
        
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = Inches(0.3)
        
        cp = ctf.paragraphs[0]
        cp.text = title
        cp.font.name = "Arial"
        cp.font.size = Pt(16)
        cp.font.bold = True
        cp.font.color.rgb = color
        
        for item in items:
            it = ctf.add_paragraph()
            it.text = f"• {item}"
            it.font.name = "Arial"
            it.font.size = Pt(12)
            it.font.color.rgb = white_text
            it.space_before = Pt(10)
            
    # -------------------------------------------------------------------------
    # SLIDE 4: LANGGRAPH WORKFLOW
    # -------------------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide)
    add_slide_header(slide, "LangGraph Agentic Swarm Workflow")
    
    # Visual grid representing the 6 sequential nodes
    grid_w = Inches(3.8)
    grid_h = Inches(2.2)
    
    nodes = [
        ("1. Research Agent", "Queries Tavily for core metadata, financials, and executives", Inches(0.5), Inches(1.8), RGBColor(99, 102, 241)),
        ("2. News Agent", "Searches for recent releases, product updates, and mergers", Inches(4.7), Inches(1.8), RGBColor(99, 102, 241)),
        ("3. SWOT Analyst", "Processes raw research & news to formulate SWOT factors", Inches(8.9), Inches(1.8), RGBColor(99, 102, 241)),
        ("4. Competitor Agent", "Evaluates peer market positioning and structural gaps", Inches(0.5), Inches(4.4), RGBColor(20, 184, 166)),
        ("5. AI Consultant", "Recommends 3 concrete AI solutions with cost, timeline, and ROI", Inches(4.7), Inches(4.4), RGBColor(20, 184, 166)),
        ("6. CEO Pitch Agent", "Generates the final summary headline and personalized CEO memo", Inches(8.9), Inches(4.4), RGBColor(20, 184, 166))
    ]
    
    for name, desc, left, top, color in nodes:
        card = slide.shapes.add_shape(1, left, top, grid_w, grid_h)
        card.fill.solid()
        card.fill.fore_color.rgb = card_color
        card.line.color.rgb = color
        card.line.width = Pt(1.5)
        
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = Inches(0.2)
        
        cp1 = ctf.paragraphs[0]
        cp1.text = name
        cp1.font.name = "Arial"
        cp1.font.size = Pt(14)
        cp1.font.bold = True
        cp1.font.color.rgb = color
        
        cp2 = ctf.add_paragraph()
        cp2.text = desc
        cp2.font.name = "Arial"
        cp2.font.size = Pt(11)
        cp2.font.color.rgb = white_text
        cp2.space_before = Pt(6)
        
    # -------------------------------------------------------------------------
    # SLIDE 5: ENGINEERING DECISIONS & CHALLENGES
    # -------------------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide)
    add_slide_header(slide, "Engineering Challenges & Production Solutions")
    
    # 4 grid blocks
    block_w = Inches(5.9)
    block_h = Inches(2.2)
    
    challenges = [
        ("LLM JSON Validation", "LLMs often wrap responses in markdown or add prefixes that break parsing.\nSolution: Custom call_llm_json regex parser that strips markdown wrappers.", Inches(0.5), Inches(1.8), primary_color),
        ("Entity Context Pollution", "Common queries like 'Apple' return mixed results from unrelated subsidiaries.\nSolution: Instructed agents to query specific terms & filter out parent noise.", Inches(6.9), Inches(1.8), primary_color),
        ("Cold-Start API Fallback", "Without Gemini or Tavily keys, the application crashes on initial startup.\nSolution: Implemented Demo Mode with built-in presets & hash generator.", Inches(0.5), Inches(4.4), accent_color),
        ("In-Memory Stream Exporter", "Writing PDF/PPT files to server disk slows response times & leaks storage.\nSolution: Direct in-memory compilation using BytesIO bytes buffers.", Inches(6.9), Inches(4.4), accent_color)
    ]
    
    for title, desc, left, top, color in challenges:
        card = slide.shapes.add_shape(1, left, top, block_w, block_h)
        card.fill.solid()
        card.fill.fore_color.rgb = card_color
        card.line.color.rgb = color
        card.line.width = Pt(1.5)
        
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = Inches(0.2)
        
        cp1 = ctf.paragraphs[0]
        cp1.text = title
        cp1.font.name = "Arial"
        cp1.font.size = Pt(14)
        cp1.font.bold = True
        cp1.font.color.rgb = color
        
        cp2 = ctf.add_paragraph()
        cp2.text = desc
        cp2.font.name = "Arial"
        cp2.font.size = Pt(11)
        cp2.font.color.rgb = white_text
        cp2.space_before = Pt(4)
        
    # -------------------------------------------------------------------------
    # SLIDE 6: CONCLUSION & THANK YOU
    # -------------------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    set_slide_bg(slide)
    
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.33), Inches(3.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "THANK YOU"
    p.font.name = "Arial"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = white_text
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "AURA BI Portfolio Presentation Complete"
    p2.font.name = "Arial"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = accent_color
    p2.space_before = Pt(8)
    p2.alignment = PP_ALIGN.CENTER
    
    p3 = tf.add_paragraph()
    p3.text = "Ready for Questions & Evaluation"
    p3.font.name = "Arial"
    p3.font.size = Pt(14)
    p3.font.color.rgb = muted_text
    p3.space_before = Pt(20)
    p3.alignment = PP_ALIGN.CENTER
    
    # Save PPTX
    prs.save(output_path)

if __name__ == "__main__":
    import sys
    output_ppt = "AURA_BI_Project_Presentation.pptx"
    if len(sys.argv) > 1:
        output_ppt = sys.argv[1]
        
    print(f"Generating professional PPT at: {output_ppt}...")
    generate_presentation_deck(output_ppt)
    print("PPT Generation complete!")
